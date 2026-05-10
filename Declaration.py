import streamlit as st
import pdfplumber
from docx import Document
import io
import re
import pytesseract
from pdf2image import convert_from_bytes
from PIL import Image

# ฟังก์ชันช่วยดึงข้อความด้วย OCR (ในกรณีที่เป็นรูปภาพ)
def extract_text_with_ocr(pdf_bytes):
    # แปลง PDF แต่ละหน้าให้เป็นรูปภาพ
    images = convert_from_bytes(pdf_bytes)
    full_text = ""
    for img in images:
        # สั่งให้ OCR อ่านภาษาไทยและอังกฤษ
        text = pytesseract.image_to_string(img, lang='tha+eng')
        full_text += text + "\n"
    return full_text

def replace_text_in_doc(doc, old_text, new_text):
    for p in doc.paragraphs:
        if old_text in p.text:
            for run in p.runs:
                run.text = run.text.replace(old_text, new_text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if old_text in p.text:
                        for run in p.runs:
                            run.text = run.text.replace(old_text, new_text)

st.title("📄 PDF OCR to Word Template")

template_file = st.file_uploader("1. อัปโหลด Word Template (.docx)", type="docx")
uploaded_pdfs = st.file_uploader("2. อัปโหลด PDF (แบบพิมพ์หรือแบบสแกน)", type="pdf", accept_multiple_files=True)

if template_file and uploaded_pdfs:
    for pdf_file in uploaded_pdfs:
        pdf_bytes = pdf_file.read()
        
        # ลองอ่านแบบ Digital Text ก่อน
        text_content = ""
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text_content += extracted + "\n"
        
        # ถ้าอ่านปกติแล้วไม่ได้ข้อความ (เป็นไฟล์สแกน) ให้ใช้ OCR
        if len(text_content.strip()) < 10:
            st.warning(f"ไฟล์ {pdf_file.name} ดูเหมือนจะเป็นรูปภาพ กำลังใช้ OCR สแกน...")
            text_content = extract_text_with_ocr(pdf_bytes)

        # ค้นหาเงินและภาษี (แบบข้ามบรรทัด)
        # ปรับ Regex ให้ยืดหยุ่นขึ้นเผื่อ OCR อ่านวรรคตอนเพี้ยน
        amount_pattern = r'จำนวนเงินที่จ่าย\s*[\n\r]*\s*([\d,]+\.?\d*)'
        tax_pattern = r'ภาษีที่หักไว้\s*[\n\r]*\s*([\d,]+\.?\d*)'
        
        amount_match = re.search(amount_pattern, text_content)
        tax_match = re.search(tax_pattern, text_content)
        
        amount_val = amount_match.group(1) if amount_match else "ไม่พบข้อมูล"
        tax_val = tax_match.group(1) if tax_match else "ไม่พบข้อมูล"
        
        # สร้างไฟล์ Word
        template_file.seek(0)
        doc = Document(template_file)
        replace_text_in_doc(doc, "XXX", amount_val)
        replace_text_in_doc(doc, "YYY", tax_val)
        
        doc_io = io.BytesIO()
        doc.save(doc_io)
        doc_io.seek(0)
        
        col1, col2 = st.columns([3, 1])
        with col1:
            st.write(f"✅ {pdf_file.name}")
            st.caption(f"เงิน: {amount_val} | ภาษี: {tax_val}")
        with col2:
            st.download_button("โหลด Word", data=doc_io, file_name=pdf_file.name.replace(".pdf", ".docx"), key=pdf_file.name)
